"""خواندن پرونده مددجویان از فایل‌های اکسل (xlsx)، ورد (docx) و متنی (txt/csv).

هر رکورد یک دیکشنری «عنوان فیلد → مقدار» است. عنوان‌های مترادف به نام استاندارد تبدیل می‌شوند
و بقیه فیلدها همان‌طور که هستند نگهداری می‌شوند."""
import csv
import io
import re

CANON = {
    "نام": ["نام", "نام و نام خانوادگی", "نام ونام خانوادگی", "نام مددجو", "مددجو", "اسم", "نام کامل", "نام و نام‌خانوادگی"],
    "کد ملی": ["کد ملی", "کدملی", "شماره ملی", "ش ملی", "کد ملي"],
    "تاریخ تولد": ["تاریخ تولد", "تولد", "سال تولد"],
    "سن": ["سن"],
    "جنسیت": ["جنسیت", "جنس"],
    "تلفن": ["تلفن", "شماره تماس", "موبایل", "شماره موبایل", "تلفن همراه", "شماره تلفن"],
    "نام ولی/سرپرست": ["ولی", "نام ولی", "سرپرست", "نام سرپرست", "ولی/سرپرست", "همراه", "نام پدر", "نام مادر"],
    "تلفن ولی": ["تلفن ولی", "شماره ولی", "تلفن سرپرست", "تماس ولی", "تلفن اضطراری", "شماره اضطراری"],
    "آدرس": ["آدرس", "نشانی", "محل سکونت"],
    "تشخیص": ["تشخیص", "نوع معلولیت", "نوع اختلال", "بیماری", "نوع بیماری", "معلولیت", "تشخیص پزشک"],
    "داروها": ["داروها", "داروهای مصرفی", "دارو"],
    "حساسیت‌ها": ["حساسیت", "حساسیت‌ها", "حساسیت ها", "آلرژی"],
    "تاریخ پذیرش": ["تاریخ پذیرش", "تاریخ ورود", "تاریخ ثبت نام", "تاریخ شروع"],
    "شماره پرونده": ["شماره پرونده", "کد پرونده", "ش پرونده", "شماره"],
    "وضعیت": ["وضعیت"],
    "توضیحات": ["توضیحات", "توضیح", "یادداشت", "ملاحظات", "شرح"],
}
_REV = {}
for canon, aliases in CANON.items():
    for a in aliases:
        _REV[a] = canon


def _clean(s) -> str:
    s = str(s if s is not None else "")
    s = s.replace("ي", "ی").replace("ك", "ک").replace("‌", " ").replace("‏", "").replace("‎", "")
    s = s.translate(str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789"))
    return re.sub(r"\s+", " ", s).strip()


def canon_key(k: str) -> str:
    k = _clean(k).strip(":：؛").strip()
    return _REV.get(k, k)


def _val(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"[ \t]+", " ", str(v)).strip()


def _from_rows(rows):
    """سطر اول سرستون، بقیه رکورد."""
    rows = [r for r in rows if any(_val(c) for c in r)]
    if not rows:
        return []
    head = [canon_key(c) if _val(c) else "" for c in rows[0]]
    out = []
    for r in rows[1:]:
        rec = {}
        for h, c in zip(head, r):
            if h and _val(c):
                rec[h] = _val(c)
        if rec:
            out.append(rec)
    return out


def _kv_pairs(rows):
    """جدول دو ستونه «عنوان | مقدار» → یک رکورد."""
    rec = {}
    for r in rows:
        cells = [_val(c) for c in r if _val(c)]
        if len(cells) >= 2:
            rec[canon_key(cells[0])] = cells[1]
    return rec


def _looks_kv(rows) -> bool:
    rows = [r for r in rows if any(_val(c) for c in r)]
    if not rows or max(len([c for c in r if _val(c)]) for r in rows) > 2:
        return False
    return len(rows) >= 2 and sum(1 for r in rows if canon_key(_val(r[0])) in CANON) >= 2


def parse_lines(text: str):
    """متن «کلید: مقدار». رکورد جدید با تکرار کلید «نام» یا خط جداکننده شروع می‌شود."""
    recs, cur, last = [], {}, None
    for raw in text.splitlines():
        line = raw.strip()
        if not line or re.fullmatch(r"[-=_*#~ ]{3,}", line):
            if re.fullmatch(r"[-=_*#~ ]{3,}", line) and cur:
                recs.append(cur)
                cur, last = {}, None
            continue
        m = re.match(r"^(.{1,40}?)\s*[:：]\s*(.*)$", line)
        if m:
            k, v = canon_key(m.group(1)), m.group(2).strip()
            if k == "نام" and "نام" in cur:
                recs.append(cur)
                cur = {}
            if v:
                cur[k] = v
            last = k
        elif last and last in cur:
            cur[last] += "\n" + line  # ادامه مقدار چندخطی
    if cur:
        recs.append(cur)
    return [r for r in recs if r]


def parse_xlsx(data: bytes):
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    out = []
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        if _looks_kv(rows):
            out.append(_kv_pairs(rows))
        else:
            out += _from_rows(rows)
    return out


def parse_docx(data: bytes):
    import docx
    d = docx.Document(io.BytesIO(data))
    out = []
    for t in d.tables:
        rows = [[c.text for c in r.cells] for r in t.rows]
        if _looks_kv(rows):
            out.append(_kv_pairs(rows))
        elif len(rows) >= 2 and any(canon_key(c) == "نام" for c in rows[0]):
            out += _from_rows(rows)
        else:
            kv = _kv_pairs(rows)
            if kv:
                out.append(kv)
    out += parse_lines("\n".join(p.text for p in d.paragraphs))
    return out


def parse_csv(text: str):
    sample = text[:2000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(io.StringIO(text), dialect))
    return _from_rows(rows)


def decode_text(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16", "cp1256", "cp1252"):
        try:
            return data.decode(enc)
        except Exception:
            continue
    return data.decode("utf-8", errors="ignore")


def parse_file(filename: str, data: bytes):
    n = (filename or "").lower()
    if n.endswith((".xlsx", ".xlsm")):
        recs = parse_xlsx(data)
    elif n.endswith(".docx"):
        recs = parse_docx(data)
    elif n.endswith(".csv"):
        recs = parse_csv(decode_text(data))
    elif n.endswith((".txt", ".md", ".text")):
        txt = decode_text(data)
        recs = parse_lines(txt)
    else:
        raise ValueError("فرمت پشتیبانی نمی‌شود (xlsx, docx, txt, csv).")
    numeric = {"کد ملی", "تلفن", "تلفن ولی", "شماره پرونده"}
    out = []
    for r in recs:
        rec = {}
        for k, v in r.items():
            k = canon_key(k)
            if v:
                rec[k] = re.sub(r"[\s\-]", "", _clean(v)) if k in numeric else v
        if rec:
            out.append(rec)
    return out


def split_name(rec: dict):
    """نام را از رکورد جدا می‌کند؛ بقیه فیلدها به‌عنوان جزئیات برمی‌گردند."""
    rec = dict(rec)
    name = rec.pop("نام", "") or ""
    if not name:
        first, last = rec.pop("نام کوچک", ""), rec.pop("نام خانوادگی", "")
        name = f"{first} {last}".strip()
    return name.strip(), rec
