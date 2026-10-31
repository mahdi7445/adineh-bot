"""ساخت فایل‌های اکسل (راست‌به‌چپ، تاریخ شمسی) برای ورود/خروج، وام، واریزی، مددجویان و امور."""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import attlogic, crypto, db, jdate

HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
HEAD_FONT = Font(bold=True, color="FFFFFF", name="Tahoma")
BODY_FONT = Font(name="Tahoma")
THIN = Side(style="thin", color="BBBBBB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
WRAP = Alignment(horizontal="right", vertical="top", wrap_text=True)


def _sheet(wb, title, headers, rows, widths=None, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title[:31]
    ws.sheet_view.rightToLeft = True
    ws.append(headers)
    for c in ws[1]:
        c.fill, c.font, c.alignment, c.border = HEAD_FILL, HEAD_FONT, CENTER, BORDER
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font, c.border = BODY_FONT, BORDER
            c.alignment = WRAP if isinstance(c.value, str) and len(c.value) > 25 else CENTER
    for i in range(1, len(headers) + 1):
        w = (widths[i - 1] if widths and i <= len(widths) else 14)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    return ws


def _bytes(wb) -> bytes:
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def _staff(staff_id=None):
    if staff_id:
        r = db.staff_by_id(staff_id)
        return [r] if r else []
    return db.roster()


def _title(s):
    return db.title_label(s)


def _st_label(sid, day):
    st = attlogic.eff_status(sid, day)
    return attlogic.STATUS[st][1] if st else "ثبت نشده"


def _reports_by(day_from, day_to):
    out = {}
    for r in db.q("SELECT * FROM reports WHERE day BETWEEN ? AND ? ORDER BY id", (day_from, day_to)):
        out.setdefault((r["staff_id"], r["day"]), []).append(crypto.dec(r["body"]))
    return out


def _detail_rows(staff, d1, d2):
    reps = _reports_by(d1, d2)
    rows = []
    for s in staff:
        recs = db.q("SELECT * FROM attendance WHERE staff_id=? AND day BETWEEN ? AND ? ORDER BY day, id", (s["id"], d1, d2))
        seen = set()
        for r in recs:
            key = (s["id"], r["day"])
            rep = "\n— ".join(reps.get(key, [])) if key not in seen else ""
            seen.add(key)
            rows.append([jdate.fmt(r["day"]), jdate.weekday(r["day"]), s["name"], _title(s), _st_label(s["id"], r["day"]),
                         r["t_in"] or "", r["t_out"] or "", jdate.hours(r["t_in"], r["t_out"]),
                         "دستی (مدیر)" if r["source"] == "admin" else "خودکار", r["note"] or "", rep])
        # روزهای بدون ساعت: وضعیت دستی (غایب/مرخصی/…) یا فقط گزارش
        extra = {pr["day"] for pr in db.q("SELECT day FROM presence WHERE staff_id=? AND day BETWEEN ? AND ?", (s["id"], d1, d2))}
        extra |= {day for (sid, day) in reps if sid == s["id"]}
        for day in sorted(extra):
            if (s["id"], day) in seen:
                continue
            rows.append([jdate.fmt(day), jdate.weekday(day), s["name"], _title(s), _st_label(s["id"], day), "", "", 0, "", "",
                         "\n— ".join(reps.get((s["id"], day), []))])
    rows.sort(key=lambda x: (x[2], x[0]))
    return rows


DETAIL_HEAD = ["تاریخ", "روز", "نام", "عنوان", "وضعیت", "ورود", "خروج", "ساعت کارکرد", "نوع ثبت", "توضیح", "گزارش فعالیت"]
DETAIL_W = [12, 10, 18, 16, 11, 8, 8, 12, 10, 22, 60]


def att_day(day, staff_id=None) -> bytes:
    wb = Workbook()
    staff = _staff(staff_id)
    rows = _detail_rows(staff, day, day)
    have = {r[2] for r in rows}
    for s in staff:
        if s["name"] not in have:
            rows.append([jdate.fmt(day), jdate.weekday(day), s["name"], _title(s), "ثبت نشده", "", "", 0, "", "", ""])
    rows.sort(key=lambda x: x[2])
    _sheet(wb, f"ورود و خروج {jdate.fmt(day).replace('/', '-')}", DETAIL_HEAD, rows, DETAIL_W, first=True)
    return _bytes(wb)


def att_month(jy, jm, staff_id=None) -> bytes:
    d1, d2 = jdate.month_bounds(jy, jm)
    n = jdate.month_days(jy, jm)
    staff = _staff(staff_id)
    wb = Workbook()
    head = ["ردیف", "نام", "عنوان"] + [str(i) for i in range(1, n + 1)] + ["جمع ساعت", "روزهای حضور", "غایب", "مرخصی", "استعلاجی", "مأموریت"]
    rows = []
    for i, s in enumerate(staff, 1):
        per_day = {}
        for r in db.q("SELECT * FROM attendance WHERE staff_id=? AND day BETWEEN ? AND ?", (s["id"], d1, d2)):
            per_day[r["day"]] = per_day.get(r["day"], 0) + jdate.hours(r["t_in"], r["t_out"])
        cells, cnt, present = [], {"absent": 0, "leave": 0, "sick": 0, "mission": 0}, 0
        for d in range(1, n + 1):
            day = jdate.from_j(jy, jm, d)
            st = attlogic.eff_status(s["id"], day)
            h = per_day.get(day)
            if st in cnt:
                cnt[st] += 1
                cells.append(attlogic.STATUS[st][2])
            else:
                cells.append(round(h, 2) if h else ("✓" if st == "present" else ""))
            present += 1 if st == "present" else 0
        rows.append([i, s["name"], _title(s)] + cells + [round(sum(per_day.values()), 2), present,
                                                          cnt["absent"], cnt["leave"], cnt["sick"], cnt["mission"]])
    ws = _sheet(wb, f"خلاصه {jdate.MONTHS[jm - 1]} {jy}", head, rows, [6, 18, 14] + [5.5] * n + [11, 12, 8, 8, 9, 9], first=True)
    # جمعه‌ها را رنگی کن
    for d in range(1, n + 1):
        if jdate.weekday(jdate.from_j(jy, jm, d)) == "جمعه":
            for r in range(1, len(rows) + 2):
                ws.cell(row=r, column=3 + d).fill = PatternFill("solid", fgColor="FCE4D6") if r > 1 else HEAD_FILL
    _sheet(wb, "جزئیات و گزارش‌ها", DETAIL_HEAD, _detail_rows(staff, d1, d2), DETAIL_W)
    return _bytes(wb)


def att_year(jy, staff_id=None) -> bytes:
    d1, d2 = jdate.year_bounds(jy)
    staff = _staff(staff_id)
    wb = Workbook()
    head = ["ردیف", "نام", "عنوان"] + jdate.MONTHS + ["جمع ساعت سال", "روزهای حضور"]
    rows = []
    for i, s in enumerate(staff, 1):
        months = [0.0] * 12
        days = set()
        for r in db.q("SELECT * FROM attendance WHERE staff_id=? AND day BETWEEN ? AND ?", (s["id"], d1, d2)):
            _, m, _ = jdate.to_j(r["day"])
            months[m - 1] += jdate.hours(r["t_in"], r["t_out"])
            days.add(r["day"])
        for pr in db.q("SELECT day FROM presence WHERE staff_id=? AND status='present' AND day BETWEEN ? AND ?", (s["id"], d1, d2)):
            days.add(pr["day"])
        rows.append([i, s["name"], _title(s)] + [round(x, 2) if x else "" for x in months] + [round(sum(months), 2), len(days)])
    _sheet(wb, f"خلاصه سال {jy}", head, rows, [6, 18, 14] + [10] * 12 + [13, 12], first=True)
    _sheet(wb, "جزئیات", DETAIL_HEAD, _detail_rows(staff, d1, d2), DETAIL_W)
    return _bytes(wb)


def loans() -> bytes:
    from .mod_money import loan_stats
    wb = Workbook()
    st = {"active": "فعال", "queued": "در صف نوبت", "done": "تسویه‌شده"}
    rows = []
    for l in db.q("SELECT * FROM loans ORDER BY status, turn_no, id"):
        paid, cnt, rem, left, nxt = loan_stats(l)
        rows.append([db.staff_name(l["staff_id"]), l["turn_no"] or "", l["amount"], l["installment"], l["n_installments"],
                     jdate.fmt(l["start_day"]), cnt, paid, rem, left, jdate.fmt(nxt) if nxt else "", st[l["status"]], l["note"] or ""])
    _sheet(wb, "وام‌ها", ["نام", "نوبت", "مبلغ وام", "قسط ماهانه", "تعداد اقساط", "شروع", "اقساط پرداخت‌شده", "جمع پرداخت",
                          "مانده", "اقساط باقی‌مانده", "سررسید بعدی", "وضعیت", "توضیح"], rows,
           [18, 7, 14, 14, 10, 12, 12, 14, 14, 12, 13, 12, 25], first=True)
    pay = []
    for p in db.q("SELECT * FROM loan_payments ORDER BY day, id"):
        l = db.one("SELECT * FROM loans WHERE id=?", (p["loan_id"],))
        pay.append([jdate.fmt(p["day"]), db.staff_name(l["staff_id"]) if l else "", p["amount"], p["note"] or ""])
    _sheet(wb, "پرداخت اقساط", ["تاریخ", "نام", "مبلغ", "توضیح"], pay, [12, 18, 14, 25])
    return _bytes(wb)


def payouts(jy) -> bytes:
    d1, d2 = jdate.year_bounds(jy)
    wb = Workbook()
    rows = [[jdate.fmt(p["day"]), db.staff_name(p["staff_id"]), p["kind"], p["amount"], p["loan_deduct"] or "", p["note"] or ""]
            for p in db.q("SELECT * FROM payouts WHERE day BETWEEN ? AND ? ORDER BY day, id", (d1, d2))]
    _sheet(wb, f"واریزی‌ها {jy}", ["تاریخ", "نام", "نوع", "مبلغ", "کسر وام", "توضیح"], rows, [12, 18, 12, 14, 12, 25], first=True)
    summ = {}
    for p in db.q("SELECT * FROM payouts WHERE day BETWEEN ? AND ?", (d1, d2)):
        _, m, _ = jdate.to_j(p["day"])
        summ.setdefault(p["staff_id"], [0] * 12)[m - 1] += p["amount"]
    srows = [[db.staff_name(sid)] + [v or "" for v in vals] + [sum(vals)] for sid, vals in summ.items()]
    _sheet(wb, "جمع ماهانه", ["نام"] + jdate.MONTHS + ["جمع سال"], srows, [18] + [12] * 12 + [14])
    return _bytes(wb)


def clients(client_id=None) -> bytes:
    wb = Workbook()
    rows = db.q("SELECT * FROM clients WHERE id=?", (client_id,)) if client_id else db.q("SELECT * FROM clients ORDER BY name")
    data = [(c, crypto.dec_json(c["data"])) for c in rows]
    keys = []
    for _, d in data:
        for k in d:
            if k not in keys:
                keys.append(k)
    head = ["شماره", "نام", "وضعیت", "تاریخ پذیرش"] + keys
    body = [[c["id"], c["name"], "فعال" if c["status"] == "active" else "ترخیص‌شده", jdate.fmt(c["admit_day"])] + [d.get(k, "") for k in keys]
            for c, d in data]
    _sheet(wb, "مددجویان", head, body, [7, 20, 12, 12] + [20] * len(keys), first=True)
    notes = []
    for n in db.q("SELECT * FROM client_notes ORDER BY client_id, id"):
        if client_id and n["client_id"] != client_id:
            continue
        c = db.one("SELECT name FROM clients WHERE id=?", (n["client_id"],))
        notes.append([c["name"] if c else "", jdate.fmt(n["day"]), db.staff_name(n["author"]), crypto.dec(n["body"])])
    _sheet(wb, "یادداشت‌ها", ["مددجو", "تاریخ", "ثبت‌کننده", "متن"], notes, [20, 12, 18, 80])
    return _bytes(wb)


def incidents(jy=None) -> bytes:
    wb = Workbook()
    sql, args = "SELECT * FROM incidents ORDER BY id", ()
    if jy:
        d1, d2 = jdate.year_bounds(jy)
        sql, args = "SELECT * FROM incidents WHERE day BETWEEN ? AND ? ORDER BY id", (d1, d2)
    sev = {"normal": "عادی", "important": "مهم", "urgent": "فوری"}
    rows = [[jdate.fmt(r["day"]), r["tm"] or "", r["title"], sev.get(r["severity"], ""), "باز" if r["status"] == "open" else "حل‌شده",
             db.staff_name(r["created_by"]), crypto.dec(r["body"]), crypto.dec(r["resolution"]), jdate.fmt(r["resolved_day"]) if r["resolved_day"] else ""]
            for r in db.q(sql, args)]
    _sheet(wb, "امور و اتفاقات", ["تاریخ", "ساعت", "عنوان", "اهمیت", "وضعیت", "ثبت‌کننده", "شرح", "نتیجه رسیدگی", "تاریخ رسیدگی"], rows,
           [12, 8, 25, 8, 10, 16, 50, 40, 12], first=True)
    return _bytes(wb)


def month_report(jy, jm, staff_id=None) -> bytes:
    """گزارش ماهانه: «کلیات» (هر نفر یک ردیف)، «روزانه» (هر روز هر نفر با ساعت ورود/خروج/کارکرد) و «گزارش‌ها» (متن کامل با نام و ساعت ثبت)."""
    staff = _staff(staff_id)
    wb = Workbook()
    month = f"{jdate.MONTHS[jm - 1]} {jy}"
    summ, daily = [], []
    for i, s in enumerate(staff, 1):
        st = attlogic.month_stats(s["id"], jy, jm)
        c = st["counts"]
        summ.append([i, s["name"], _title(s), c["present"], c["absent"], c["leave"], c["sick"], c["mission"], c["unrecorded"],
                     c["incomplete"], st["worked_days"], st["hours"], jdate.hm(st["hours"]), jdate.hm(st["avg"]), st["reports"]])
        for d in st["days"]:
            if d["future"]:
                continue
            label = attlogic.STATUS[d["status"]][1] if d["status"] else ("جمعه" if d["friday"] else "ثبت نشده")
            recs = d["recs"] or [None]
            for k, r in enumerate(recs):
                daily.append([jdate.fmt(d["day"]), jdate.weekday(d["day"]), s["name"], _title(s), label,
                              (r["t_in"] or "") if r else "", (r["t_out"] or "") if r else "",
                              jdate.hours(r["t_in"], r["t_out"]) if r else 0, d["hours"] if k == 0 else "",
                              (r["note"] or "") if r else "", len(d["reports"]) if k == 0 else ""])
    head = ["ردیف", "نام", "عنوان", "روز حضور", "غیبت", "مرخصی", "استعلاجی", "مأموریت", "بدون وضعیت", "ساعت ناقص",
            "روزهای دارای ساعت", "جمع ساعت (عدد)", "جمع ساعت (ساعت:دقیقه)", "میانگین روزانه", "تعداد گزارش"]
    if len(summ) > 1:
        tot = ["", "جمع کل", ""] + [sum(r[k] for r in summ) for k in range(3, 12)] + [jdate.hm(sum(r[11] for r in summ)), "", sum(r[14] for r in summ)]
        summ.append(tot)
    _sheet(wb, f"کلیات {month}"[:31], head, summ, [6, 20, 16] + [10] * 9 + [14, 16, 14, 12], first=True)
    _sheet(wb, "روزانه (جزئیات)", ["تاریخ", "روز", "نام", "عنوان", "وضعیت", "ورود", "خروج", "ساعت این ثبت", "جمع ساعت روز", "توضیح", "تعداد گزارش"],
           daily, [12, 10, 20, 16, 12, 8, 8, 12, 12, 24, 12])
    d1, d2 = jdate.month_bounds(jy, jm)
    reps = []
    for s in staff:
        for r in db.q("SELECT * FROM reports WHERE staff_id=? AND day BETWEEN ? AND ? ORDER BY day, id", (s["id"], d1, d2)):
            reps.append([jdate.fmt(r["day"]), jdate.weekday(r["day"]), jdate.local_hhmm(r["created"] or ""), s["name"], _title(s), crypto.dec(r["body"])])
    _sheet(wb, "گزارش‌های فعالیت", ["تاریخ", "روز", "ساعت ثبت", "نام", "عنوان", "متن گزارش"], reps, [12, 10, 10, 20, 16, 90])
    return _bytes(wb)
